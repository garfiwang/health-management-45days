import os
import pptx
from pptx.util import Inches

def update_pptx_slide3():
    pptx_path = "健管45天圖片簡報-王執定.pptx"
    slide03_img = "generated/slide03_richman_crisis_preview.png"

    if not os.path.exists(pptx_path):
        print("Missing pptx")
        return

    prs = pptx.Presentation(pptx_path)
    
    # 如果已經有 3 頁，移除第 3 頁再重新加，或直接修改
    # 在 python-pptx 中最乾淨的方式是重建投影片，或替換 shape
    # 由於我們只有 3 頁，直接重建 3 頁確保最乾淨無冗餘
    new_prs = pptx.Presentation()
    new_prs.slide_width = Inches(13.333)
    new_prs.slide_height = Inches(7.5)
    blank_layout = new_prs.slide_layouts[6]

    # Slide 1
    s1 = new_prs.slides.add_slide(blank_layout)
    s1.shapes.add_picture("generated/slide01_richman_ip_16_9_perfect.png", 0, 0, width=new_prs.slide_width, height=new_prs.slide_height)
    s1.notes_slide.notes_text_frame.text = (
        "【開場｜預計 30 秒】\n"
        "大家好，我是 Rich。\n"
        "今天要來和大家分享 45 天的蛻變歷程。"
    )

    # Slide 2
    s2 = new_prs.slides.add_slide(blank_layout)
    s2.shapes.add_picture("generated/slide02_richman_ordinary_world_preview.png", 0, 0, width=new_prs.slide_width, height=new_prs.slide_height)
    s2.notes_slide.notes_text_frame.text = (
        "【第一幕・平凡世界｜預計 60 秒】\n\n"
        "別人都是朝九晚五，而我平常的工作是『朝五晚九』。\n\n"
        "什麼叫做朝五晚九呢？大家有聽過 BNI 商會嗎？應該有聽過吧。BNI 是不是都要很早起床？沒有錯。\n\n"
        "我有兩個 BNI 商會，一個是我自己創立的長輝分會，另一個是我輔導的金佑分會。所以我平常都很早起來開會，開完會之後要整理財經資料，下午要拜訪客戶或是接受客戶諮詢，有的時候會有演講，晚上也常常會有應酬和飯局。\n\n"
        "所以，可能 5 點就要起床，一路忙到晚上九點。\n\n"
        "剛退伍的時候，可以拜訪客戶到晚上 12 點，第二天早上 8 點一樣準時進公司開會。\n"
        "到了 30 歲以後，早上就會需要喝一杯咖啡來提神。\n"
        "到了 40 歲以後，喝一杯咖啡可能不夠了，要喝到三杯咖啡，還會睡過頭爬不起來。"
    )

    # Slide 3
    s3 = new_prs.slides.add_slide(blank_layout)
    s3.shapes.add_picture(slide03_img, 0, 0, width=new_prs.slide_width, height=new_prs.slide_height)
    s3.notes_slide.notes_text_frame.text = (
        "【第二幕・冒險召喚｜預計 75 秒】\n\n"
        "這時候你就會常常聽到新聞報導很多名人突然猝死，有的是因為過勞，最近也有很多是他平時不抽菸、不喝酒、很養生，但還是會出事情。\n"
        "聽到了這些新聞，會不會讓你覺得：『我應該要休息一下？我應該要來做做健管？』\n"
        "不會～～因為我覺得我的身體很好呀！我家的基因很健康，我平常都有做好事、累積福報，所以我不會像他們這些人一樣！\n\n"
        "後來我的身體給我了一個警告：晚上睡覺半夜，膝蓋突然劇烈痛了起來。\n"
        "那種痛，會讓你痛到發瘋！請問大家，你覺得是什麼毛病？\n"
        "（等待台下回答：痛風！）\n"
        "痛風？！不可能！我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我怎麼可能有痛風？！\n\n"
        "後來我的身體，再給了我第二個警告。\n"
        "因為我的腰圍越來越寬廣，我就想說應該要來執行 168 斷食，所以我不吃早餐，早上只喝黑咖啡，想說這樣可以燃燒脂肪。\n"
        "結果有一天早上，突然頭暈目眩、心跳加速，整個身體完全提不上力，趕快去醫院檢查，結果一量血壓——直接飆到 160！\n\n"
        "到了這個時候，你覺得是不是應該要讓身體休息一下，要來做做健管，重新讓身體修復了？\n"
        "不用啦～～我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我多喝水，多上健身房運動就可以了啦！"
    )

    new_prs.save(pptx_path)
    print(f"Updated {pptx_path} with 3 pristine slides.")

if __name__ == "__main__":
    update_pptx_slide3()
