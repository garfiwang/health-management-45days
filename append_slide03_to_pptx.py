import os
import pptx
from pptx.util import Inches

def update_presentation_slide3():
    pptx_path = "健管45天圖片簡報-王執定.pptx"
    slide03_img = "generated/slide03_richman_crisis_preview.png"

    if not os.path.exists(pptx_path):
        print(f"Error: {pptx_path} not found")
        return

    prs = pptx.Presentation(pptx_path)

    blank_layout = prs.slide_layouts[6]
    slide3 = prs.slides.add_slide(blank_layout)

    # 加入第 3 頁全圖（16:9 無變形，13.333 x 7.5 英吋）
    if os.path.exists(slide03_img):
        slide3.shapes.add_picture(slide03_img, 0, 0, width=prs.slide_width, height=prs.slide_height)

    # 加入第 3 頁演講者備忘講稿（Rich 親撰，含痛風互動與積福報自嘲笑點）
    slide3.notes_slide.notes_text_frame.text = (
        "【第二幕・冒險召喚｜預計 75 秒】\n\n"
        "這時候你就會常常聽到新聞報導很多名人突然猝死，有的是因為過勞，最近也有很多是他平時不抽菸、不喝酒、很養生，但還是會出事情。\n"
        "聽到了這些新聞，會不會讓你覺得：『我應該要休息一下？我應該要來做做健管？』\n"
        "不會～～因為我覺得我的身體很好呀！我家的基因很健康，我平常都有做好事、累積福報，所以我不會像他們這些人一樣！\n\n"
        "後來我的身體給我了一個警告：晚上睡覺半夜，膝蓋突然劇烈痛了起來。\n"
        "那種痛，會讓你痛到發瘋！請問大家，你覺得是什麼毛病？\n"
        "（等待台下回答：痛風！）\n"
        "痛風？！不可能！我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我怎麼可能有痛風？！\n\n"
        "後來我的身體，再給了我第二個警告。\n"
        "因為我的腰圍越來越寬廣，我就想說應該要來執行 168 斷食，所以我不吃早餐，早上只喝黑咖啡，想說這樣可以燃燒脂肪。\n"
        "結果有一天早上，突然頭暈目眩、心跳加速，整個身體完全提不上力，趕快去醫院檢查，結果一量血壓——直接飆到 160！\n\n"
        "到了這個時候，你覺得是不是應該要讓身體休息一下，要來做做健管，重新讓身體修復了？\n"
        "不用啦～～我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我多喝水，多上健身房運動就可以了啦！"
    )

    prs.save(pptx_path)
    print(f"Successfully added Slide 3 to {pptx_path}. Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    update_presentation_slide3()
