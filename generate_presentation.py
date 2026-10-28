import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Nutrilite Style Background Images
IMG_COVER = "generated/nutrilite_cover_sunrise_20261029_004204.png"
IMG_CONTENT = "generated/nutrilite_content_clean_20261029_004229.png"
IMG_CRISIS = "generated/nutrilite_crisis_breakthrough_20261029_004253.png"

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    C_DARK_BG = RGBColor(15, 23, 42)       # Slate 900
    C_LIGHT_BG = RGBColor(248, 250, 252)   # Slate 50
    C_CARD_BG = RGBColor(255, 255, 255)    # Pure White
    C_CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    C_PRIMARY = RGBColor(5, 150, 105)      # Emerald 600
    C_PRIMARY_LIGHT = RGBColor(209, 250, 229) # Emerald 100
    C_AMBER = RGBColor(217, 119, 6)        # Amber 600
    C_RED = RGBColor(220, 38, 38)          # Red 600
    C_TEXT_DARK = RGBColor(15, 23, 42)     # Slate 900
    C_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    C_TEXT_LIGHT = RGBColor(255, 255, 255) # White

    def set_slide_background(slide, image_path, fallback_color=C_LIGHT_BG):
        if os.path.exists(image_path):
            slide.shapes.add_picture(image_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        else:
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = fallback_color
            bg.line.fill.background()

    def add_header(slide, step_num, step_name, title, subtitle=""):
        # Header tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = f"HERO'S JOURNEY ｜ 第 {step_num} 幕・{step_name}"
        p_tag.font.name = "Arial"
        p_tag.font.size = Pt(13)
        p_tag.font.bold = True
        p_tag.font.color.rgb = C_PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Microsoft JhengHei"
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_DARK

        if subtitle:
            p_sub = tf_title.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Microsoft JhengHei"
            p_sub.font.size = Pt(14)
            p_sub.font.color.rgb = C_TEXT_MUTED

    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, IMG_COVER, C_DARK_BG)

    # Dark Frosted Container Card for Title
    add_card(slide1, Inches(0.8), Inches(1.3), Inches(11.733), Inches(4.9), bg_color=RGBColor(15, 23, 42), border_color=C_PRIMARY)

    # Accent decorative bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.15), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_PRIMARY
    bar.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.6), Inches(1.8), Inches(10.5), Inches(3.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "健管 45 天的心得與蛻變 ｜ 紐崔萊自然科學實踐"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(18)
    p.font.color.rgb = C_PRIMARY
    p.font.bold = True
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "重啟 48 歲的身體主控權"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_LIGHT
    p.space_after = Pt(12)

    p = tf1.add_paragraph()
    p.text = "—— 一個上班族的「英雄旅程」上台分享"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_after = Pt(26)

    p = tf1.add_paragraph()
    p.text = "分享者：王執定（Rich） ｜ 45天減重 5.7kg・腰圍 -8cm・身體年齡 -6歲"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(203, 213, 225)

    slide1.notes_slide.notes_text_frame.text = (
        "【開場｜預計 30 秒】\n"
        "大家好，我是 Rich。\n"
        "今天要來和大家分享 45 天的蛻變歷程。"
    )

    # ==========================================
    # SLIDE 2: 平凡世界
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide2, "一", "平凡世界", "朝五晚九的日常：不知不覺走上的「咖啡滑坡」", "剛退伍能拚到半夜，如今喝三杯咖啡還爬不起來")

    card_w = Inches(3.64)
    card_h = Inches(4.3)
    c1 = add_card(slide2, Inches(0.8), Inches(2.2), card_w, card_h)
    c2 = add_card(slide2, Inches(4.84), Inches(2.2), card_w, card_h)
    c3 = add_card(slide2, Inches(8.88), Inches(2.2), card_w, card_h)

    # Card 1 text
    tb = slide2.shapes.add_textbox(Inches(1.1), Inches(2.5), Inches(3.04), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "01. 朝五晚九的節奏"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 創立 BNI 長輝分會、輔導金佑分會\n• 清晨 5 點早起開會，白天整理財經資料\n• 下午拜訪客戶諮詢、演講\n• 晚上應酬飯局，天天忙到 9 點"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    # Card 2 text
    tb = slide2.shapes.add_textbox(Inches(5.14), Inches(2.5), Inches(3.04), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "02. 體力的殘酷滑坡"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_AMBER
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 退伍時：拚到半夜 12 點，隔天 8 點準時開會\n• 30 歲：早上需要 1 杯咖啡提神\n• 40 歲：喝到 3 杯咖啡還不夠\n• 甚至睡過頭、早上爬不起來！"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    # Card 3 text
    tb = slide2.shapes.add_textbox(Inches(9.18), Inches(2.5), Inches(3.04), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "03. 鐵齒第一心態"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_TEXT_MUTED
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 聽見名人過勞猝死新聞...\n• 心想：『不用休息啦！』\n• 『我身體好、我家基因那麼健康』\n• 『我平常都有做好事、積福報！』"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    slide2.notes_slide.notes_text_frame.text = (
        "【第一幕・平凡世界｜預計 60 秒】\n"
        "別人都是朝九晚五，而我平常的工作是「朝五晚九」。\n"
        "什麼叫做朝五晚九呢？大家有聽過 BNI 商會嗎？應該有聽過吧。BNI 是不是都要很早起床？沒有錯。\n"
        "我有兩個 BNI 商會，一個是我自己創立的長輝分會，另一個是我輔導的金佑分會。所以我平常都很早起來開會，開完會之後要整理財經資料，下午要拜訪客戶或是接受客戶諮詢，有的時候會有演講，晚上也常常會有應酬和飯局。\n"
        "所以 可能5點要起床，忙到晚上九點。\n"
        "剛退伍的時候，可以拜訪客戶到晚上 12 點，第二天早上 8 點一樣準時進公司開會。\n"
        "到了 30 歲以後，早上就會需要喝一杯咖啡來提神。\n"
        "到了 40 歲以後，喝一杯咖啡可能不夠了，要喝到三杯咖啡，還會睡過頭爬不起來。"
    )

    # ==========================================
    # SLIDE 3: 冒險召喚
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide3, "二", "冒險召喚", "身體的兩記重擊：痛風發作與血壓飆破 160", "當「積福報」擋不住身體求救，才發現盲目 168 是玩命")

    add_card(slide3, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3))
    add_card(slide3, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.3))

    tb = slide3.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(4.8), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ 警告一：痛風半夜突襲"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_RED
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 警訊爆發：睡覺半夜膝蓋突然劇痛，痛到發瘋！\n\n• 答案揭曉：『痛風』！\n\n• 內心震撼與鐵齒：\n  『不可能！我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我怎麼可能有痛風？！』"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    tb = slide3.shapes.add_textbox(Inches(7.2), Inches(2.6), Inches(4.9), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ 警告二：盲目 168 送急診"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_AMBER
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 腰圍變寬想燃脂：早餐不吃，猛灌黑咖啡！\n\n• 慘劇發生：頭暈目眩、心悸、全身提不上力！\n\n• 送醫一量：血壓直接飆破 160！\n\n• 依然鐵齒：\n  『不用健管啦！我身體好、基因好、有積福報，多喝水運動就好了啦！』"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    slide3.notes_slide.notes_text_frame.text = (
        "【第二幕・冒險召喚｜預計 75 秒】\n"
        "這時候你就會常常聽到新聞報導很多名人突然猝死，有的是因為過勞，最近也有很多是他平時不抽菸、不喝酒、很養生，但還是會出事情。聽到了這些新聞，會不會讓你覺得我應該要休息一下？我應該要來做做健管。\n"
        "不會~~因為我覺得我的身體很好呀，我家的基因很健康，我平常都有做好事、累積福報，所以我不會像他們這些人一樣。\n"
        "後來我的身體給我了一個警告：晚上睡覺半夜膝蓋會突然痛起來。\n"
        "那種痛會讓你痛到發瘋。請問大家，你覺得是什麼毛病？ 痛風\n"
        "不可能，我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我怎麼可能有痛風？\n"
        "後來我的身體再給了我一個警告。\n"
        "因為我的腰圍越來越寬廣，我就想說應該要來執行 168，所以我不吃早餐，早上就喝黑咖啡，想說可以燃燒脂肪。\n"
        "結果有一天早上頭暈目眩、心跳加速，整個身體提不上力，趕快去醫院檢查，結果一量血壓飆到 160\n"
        "到了這個時候，你覺得是不是應該要讓身體休息一下，要來做做健管，重新讓身體修復？\n"
        "不用啦~~我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我多喝水，多上健身房運動就可以了啦"
    )

    # ==========================================
    # SLIDE 4: 跨越門檻
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide4, "三", "跨越門檻", "簽下 45 天承諾契約：告別舒適圈，這一次不留退路", "跨越門檻最難的不是拒絕美食，而是下定決心「不找藉口」")

    add_card(slide4, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.3))

    tb = slide4.shapes.add_textbox(Inches(1.3), Inches(2.6), Inches(10.7), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "踏出舒適圈的三大「斷捨離」"
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(20)

    p = tf.add_paragraph()
    p.text = "1. 清空廚房與冰箱：把加工甜食、含糖飲料、宵夜零嘴通通送走，不給自己動搖的機會。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "2. 簽訂承諾契約：訂下目標體重 65.27kg（-3.44kg）、體脂率 -20%，在教練面前正式宣誓立約。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "3. 建立全新紀律：每天喝足 3000cc 以上好水、落實優質蛋白日與纖體日，準時記錄上傳餐盤。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK

    slide4.notes_slide.notes_text_frame.text = (
        "【第三幕・跨越門檻｜預計 45 秒】\n"
        "英雄旅程的第三步，是跨越門檻——踏出那一步，就代表舊的世界結束，沒有退路了。\n"
        "2026 年 7 月 30 日，我和 Carolyn 一起正式簽下了健管 45 天的承諾書。\n"
        "我給自己設下的目標是減重 3.44 公斤（-5%），體脂下降 20%。\n"
        "我們做了一件很徹底的事：把家裡的零食、含糖飲料全部清空。我們不再依賴以前的高油高鹽外食，而是每天乖乖備齊優質蛋白質、大量蔬菜，每天喝下至少 3000cc 的水。\n"
        "跨越門檻的那天，我知道這不再是『試試看』，而是一場對自己的承諾。"
    )

    # ==========================================
    # SLIDE 5: 試煉與盟友
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide5, "四", "試煉與盟友", "一個人走得快，兩個人走得遠：盟友與教練的力量", "在嚴格的試煉中，是同儕的信任與專業導航讓我們堅持向前")

    add_card(slide5, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3))
    add_card(slide5, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.3))

    tb = slide5.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(4.8), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👥 最強革命戰友：Carolyn"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 廚房裡的默契：一起採買、秤重、準備 211 纖體餐盤與蛋白日餐點。\n\n• 互相踢館與激勵：當一個人想偷懶放鬆時，另一個人就是最嚴格的糾察隊。\n\n• 雙人共好成果：45 天下來，Carolyn 純脂肪減掉 1.6 kg（佔減重 76%），肌率逆勢上升！"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    tb = slide5.shapes.add_textbox(Inches(7.2), Inches(2.6), Inches(4.9), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🧭 智慧導航者：麥淑媛教練與沛璇"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_AMBER
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "• 科學方法導入：打破『少吃就會瘦』的迷思，教我們『吃對順序、吃足營養』。\n\n• 即時盤查與診斷：每一餐拍照回報，教練一眼看出油脂超標或水分不足。\n\n• 心靈定海神針：在情緒波動與身體適應期，給予最強大的心理支持。"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    slide5.notes_slide.notes_text_frame.text = (
        "【第四幕・試煉與盟友｜預計 55 秒】\n"
        "任何英雄旅程都不可能單打獨鬥。\n"
        "在這段旅程中，我最感謝的盟友就是我的另一半 Carolyn。\n"
        "俗話說：『一個人走得快，兩個人走得遠。』\n"
        "我們兩個人一起在廚房備餐、一起討論蛋白日的菜單、互相叮嚀喝水進度。當其中一個人感到嘴饞想偷吃時，另一個人就是最堅固的防線！最後 Carolyn 也減掉了 1.6 公斤純脂肪，骨骼肌不減反增！\n"
        "更重要的是我們的導航者——麥淑媛教練與沛璇。\n"
        "她們不是只給我們菜單，而是每天在群組看我們的餐盤照片，教我們辨認隱形地雷，讓我們明白原來吃對食物、攝取足夠優質蛋白質，身體真的會以驚人的速度修復。"
    )

    # ==========================================
    # SLIDE 6: 最大危機（一）：肌耐力大考驗
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, IMG_CRISIS, C_LIGHT_BG)
    add_header(slide6, "五", "最大危機（一）", "深淵試煉：爬 5 樓雙腿發軟，我是不是把肌肉減掉了？", "體重神速下降的背後，迎來了對體能與心理的最大衝擊")

    add_card(slide6, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.3))

    tb = slide6.shapes.add_textbox(Inches(1.3), Inches(2.6), Inches(10.7), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "第 3 週的撞牆期：雙腿發軟的恐慌"
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = C_RED
    p.space_after = Pt(18)

    p = tf.add_paragraph()
    p.text = "• 體重大幅減輕的驚喜：第 3 週體重已經降到 63.6 kg（提前超越目標，達標率 148%）！"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• 意想不到的體能危機：有一天爬 5 樓時，突然感到雙腿沉重、肌耐力嚴重不足，心跳劇烈狂跳。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_RED
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• 內心的懷疑與心魔：『我是不是減太快了？我是不是把寶貴的肌肉給減掉了？這樣真的健康嗎？』"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• 教練的科學破解：這是身體燃脂途徑轉換與醣原排空後的正常適應期！立即補足電解質、優化蛋白質補充，危機迎刃而解。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_PRIMARY

    slide6.notes_slide.notes_text_frame.text = (
        "【第五幕・最大危機（一）｜預計 55 秒】\n"
        "接下來，來到了英雄旅程最險峻的深淵——最大危機。\n"
        "大家看我的數據，第一週就狂瘦 3.2 公斤，到第三週體重掉到 63.6 公斤，達標率高達 148%。\n"
        "但就在大家為我鼓掌的時候，我迎來了第一個巨大考驗。\n"
        "有一天我走樓梯爬到 5 樓，雙腿突然發軟無力，整個人喘到不行。\n"
        "那一瞬間，我心裡冒出了極大的恐懼與懷疑：『我是不是做錯了？我是不是只減掉了水分跟肌肉？這樣虛弱的瘦，真的是我要的嗎？』\n"
        "我立刻向淑媛教練求助。教練看過我的數據與狀態，耐心地安撫我：這是大量燃脂時體內水分與肝醣快速消耗的必經轉型期，只要調整營養補給與水分電解質，身體很快就會適應。果不其然，調整後體力迅速回升，打破了我的心魔！"
    )

    # ==========================================
    # SLIDE 7: 最大危機（二）：素食齋戒週
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, IMG_CRISIS, C_LIGHT_BG)
    add_header(slide7, "五", "最大危機（二）", "自律極限：素食齋戒週，在無肉少油中的終極拔河", "當習慣的蛋白質來源全部歸零，靠什麼撐過這場身心風暴？")

    add_card(slide7, Inches(0.8), Inches(2.2), Inches(11.7), Inches(4.3))

    tb = slide7.shapes.add_textbox(Inches(1.3), Inches(2.6), Inches(10.7), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "意志力的終極考驗：素食齋戒挑戰"
    p.font.bold = True
    p.font.size = Pt(24)
    p.font.color.rgb = C_RED
    p.space_after = Pt(18)

    p = tf.add_paragraph()
    p.text = "• 蛋白質攝取的巨大斷層：平時習慣吃雞胸肉、魚肉、牛肉，進入齋戒週後所有肉類葷食全部歸零！"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• 外食地雷處處驚心：市售素食往往重油、勾芡、加工素料充斥，要找到無油、純淨的植物蛋白難如登天。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• 飢餓感與自律拔河：面對社交場合與撲鼻的美食香氣，必須依靠純植物性蛋白質補充與極致自律。"
    p.font.size = Pt(17)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "• 頓悟與轉化：這不是『剝奪』，而是一次給腸胃道最極致的深層大掃除與排毒淨化！"
    p.font.size = Pt(17)
    p.font.color.rgb = C_PRIMARY

    slide7.notes_slide.notes_text_frame.text = (
        "【第五幕・最大危機（二）｜預計 55 秒】\n"
        "如果說肌耐力是體能的考驗，那『素食齋戒週』就是心智與自律的最高峰。\n"
        "在健管過程中碰上齋戒期，平常習慣依靠的雞肉、魚肉、蛋等高蛋白食物全部不能碰。\n"
        "走在外面，台灣傳統素食店幾乎都是炸豆皮、素肉、勾芡和滿滿的沙拉油，要做到健管要求的『純淨、無油、高蛋白』，簡直寸步難行！\n"
        "那幾天，飢餓感與想要破戒的誘惑無時無刻在腦中交戰。\n"
        "我和 Carolyn 只能想盡辦法，用植物蛋白素、大量深色蔬菜和豆腐來嚴格維持營養。\n"
        "當我們咬緊牙關撐過那週後，奇蹟發生了——我的腸道排便變得極度順暢純淨，身體前所未有的輕盈！那一刻我領悟到：限制不是剝奪，限制是為了讓身體找回原本強大的自癒能力。"
    )

    # ==========================================
    # SLIDE 8: 蛻變與寶物（數據篇）
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide8, "六", "蛻變與寶物（數據篇）", "45 天真實成績單：數字不會說謊，身體給了最誠實的回饋", "體態脫胎換骨，逆勢增肌減脂！")

    # 4 Stat cards
    stats = [
        ("體重變化", "68.7 ➔ 63.0", "累減 5.7 kg", "達標率 165.7%", C_PRIMARY),
        ("腰圍尺寸", "90 ➔ 82 cm", "狂瘦 8 公分", "皮帶多退 3 格", C_PRIMARY),
        ("身體年齡", "48 ➔ 42 歲", "年輕 6 歲", "逆齡生長！", C_AMBER),
        ("骨骼肌率", "31.1% ➔ 31.7%", "逆勢上升 +0.6%", "減脂不減肌", C_PRIMARY),
    ]

    for i, (stitle, snum, sdiff, ssub, scolor) in enumerate(stats):
        c_left = Inches(0.8 + i * 2.98)
        card = add_card(slide8, c_left, Inches(2.2), Inches(2.8), Inches(4.3))
        tb = slide8.shapes.add_textbox(c_left + Inches(0.2), Inches(2.5), Inches(2.4), Inches(3.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_MUTED
        p.space_after = Pt(14)

        p = tf.add_paragraph()
        p.text = snum
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = sdiff
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = scolor
        p.space_after = Pt(10)

        p = tf.add_paragraph()
        p.text = ssub
        p.font.size = Pt(14)
        p.font.color.rgb = C_TEXT_MUTED

    slide8.notes_slide.notes_text_frame.text = (
        "【第六幕・蛻變與寶物（數據篇）｜預計 55 秒】\n"
        "走過考驗，我們迎來了英雄的寶物。\n"
        "這是 45 天後我的真實成績單：\n"
        "體重：從 68.7 公斤降到 63.0 公斤，整整減掉 5.7 公斤，達成當初目標的 165.7%！\n"
        "腰圍：從 90 公分縮小到 82 公分，整整少了 8 公分！家裡所有的褲子都變鬆了，皮帶重新打了三個洞！\n"
        "身體年齡：從 48 歲回春到 42 歲，年輕了 6 歲！\n"
        "內臟脂肪：從危險的 9 降到標準的 7！\n"
        "血壓：從 140/100 降回完全健康的 120/80！\n"
        "最讓我自豪的是最後一項——我的骨骼肌率從 31.1% 逆勢上升到 31.7%！這證明了透過科學健管，我們減掉的全是多餘的脂肪與廢水，肌肉扎扎實實地被保護了下來！"
    )

    # ==========================================
    # SLIDE 9: 蛻變與寶物（生活篇）
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide9, "六", "蛻變與寶物（心態篇）", "比體重機數字更珍貴的，是「感官與精力的重生」", "這不只是外表的改變，而是內在生命力的全面升級")

    add_card(slide9, Inches(0.8), Inches(2.2), Inches(3.64), Inches(4.3))
    add_card(slide9, Inches(4.84), Inches(2.2), Inches(3.64), Inches(4.3))
    add_card(slide9, Inches(8.88), Inches(2.2), Inches(3.64), Inches(4.3))

    # Box 1
    tb = slide9.shapes.add_textbox(Inches(1.1), Inches(2.5), Inches(3.04), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👅 味蕾的重新覺醒"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(14)
    p = tf.add_paragraph()
    p.text = "• 戒斷加工糖與重鹹後，味蕾變得極度敏銳。\n• 重新品嚐出一碗清蒸高麗菜的鮮甜。\n• 不再被化學調味料綁架，真正享受食物的原貌。"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    # Box 2
    tb = slide9.shapes.add_textbox(Inches(5.14), Inches(2.5), Inches(3.04), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "⚡ 下午精力滿格不宕機"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_AMBER
    p.space_after = Pt(14)
    p = tf.add_paragraph()
    p.text = "• 下午兩點開會不再陷入昏睡絕境。\n• 告別咖啡續命，專注力與工作效率倍增。\n• 晚上睡眠品質大幅提升，起床神清氣爽。"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    # Box 3
    tb = slide9.shapes.add_textbox(Inches(9.18), Inches(2.5), Inches(3.04), Inches(3.7))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "👑 拿回人生主控權"
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(14)
    p = tf.add_paragraph()
    p.text = "• 從『被慾望控制』轉變成『主動選擇健康』。\n• 明白自己吃下肚的每一口食物會帶來什麼影響。\n• 建立終身可持續的飲食與生活新常態。"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    slide9.notes_slide.notes_text_frame.text = (
        "【第六幕・蛻變與寶物（生活篇）｜預計 50 秒】\n"
        "除了體重計上的數字，這 45 天帶給我最深沉的寶物是什麼？\n"
        "第一，是我的味蕾甦醒了。以前無辣不歡、重油重鹹；現在一口清蒸蔬菜，我能吃出大地原本的甘甜，身體會自然渴望乾淨的食物。\n"
        "第二，是精力的徹底解放。那種吃完午餐不再昏睡、下午腦袋清晰如鏡的專注感，是再多咖啡都換不來的。\n"
        "第三，也是最重要的——我重新贏回了對自己的掌控感。\n"
        "當你看見自己的身體在自己的自律與科學管理下，一步一步變成精實、健康的模樣，那種由內而外的自信，會擴散到你的工作、你的事業與你的人際關係中。"
    )

    # ==========================================
    # SLIDE 10: 帶著禮物歸來
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, IMG_CONTENT, C_LIGHT_BG)
    add_header(slide10, "七", "帶著禮物歸來", "把禮物送給現場的你：明天就能啟動的健康行動", "你不需要很厲害才開始，但你需要開始才會變得很厲害")

    add_card(slide10, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.3))
    add_card(slide10, Inches(6.8), Inches(2.2), Inches(5.7), Inches(4.3))

    tb = slide10.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(4.8), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🌱 給還沒接觸過健管的朋友"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_PRIMARY
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "明天起床，請立刻試試這兩個「微行動」：\n\n1. 【喝足體重×40cc 的好水】：\n   60公斤喝 2400cc，水是代謝脂肪最好的燃燒劑！\n\n2. 【改變進食順序：水➔肉➔菜➔飯】：\n   先吃蛋白質與蔬菜增加飽足感，血糖不震盪，下午就不昏睡！"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    tb = slide10.shapes.add_textbox(Inches(7.2), Inches(2.6), Inches(4.9), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🔥 給同行學員與教練團隊"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = C_AMBER
    p.space_after = Pt(16)
    p = tf.add_paragraph()
    p.text = "結案不是終點，而是新起點：\n\n1. 【維持期才是真功夫】：\n   健管 45 不是短暫的節食減肥，而是一輩子可長可久的健康生活哲學。\n\n2. 【成為別人的燈塔】：\n   用我們身上看得到的改變與光彩，去影響身邊深愛的家人與朋友！"
    p.font.size = Pt(16)
    p.font.color.rgb = C_TEXT_DARK

    slide10.notes_slide.notes_text_frame.text = (
        "【第七幕・帶著禮物歸來｜預計 55 秒】\n"
        "英雄旅程的終點，是把在冒險中淬鍊出的寶物帶回平凡世界，分享給身邊的人。\n"
        "今天在座的朋友，如果你還沒參加過健管，甚至覺得減肥很痛苦、很遙遠。\n"
        "我想送你兩個明天立刻能做的小禮物：第一，明天開始把水喝足（體重乘 40cc）；第二，吃飯時先吃肉跟菜，把澱粉留到最後。只要這兩個小改變，你就會發現身體有不可思議的回應。\n"
        "而對於跟我一起並肩作戰的同期夥伴與教練們：\n"
        "45 天結束了，但我們的新人生才剛開始。維持期才是我們展現實力的時刻。讓我們把這份健康的習慣維持一輩子，並成為身邊親友的健康榜樣！"
    )

    # ==========================================
    # SLIDE 11: 英雄謝幕
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, IMG_COVER, C_DARK_BG)

    # Dark Frosted Container Card for Closing
    add_card(slide11, Inches(0.8), Inches(1.3), Inches(11.733), Inches(4.9), bg_color=RGBColor(15, 23, 42), border_color=C_PRIMARY)

    bar2 = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.15), Inches(3.8))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = C_PRIMARY
    bar2.line.fill.background()

    tb11 = slide11.shapes.add_textbox(Inches(1.6), Inches(1.8), Inches(10.5), Inches(3.8))
    tf11 = tb11.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "48 歲，我的人生下半場才剛開始！"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(20)
    p.font.color.rgb = C_PRIMARY
    p.font.bold = True
    p.space_after = Pt(14)

    p = tf11.add_paragraph()
    p.text = "「什麼事情都可以放下，唯獨不能放下筷子；\n 什麼事情都可以失去，唯獨不能失去健康。」"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_AMBER
    p.space_after = Pt(18)

    p = tf11.add_paragraph()
    p.text = "45 天前，我以為我只是在管理體重；45 天後，我發現我贏回了人生的主控權。"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(20)
    p.font.color.rgb = C_TEXT_LIGHT
    p.space_after = Pt(24)

    p = tf11.add_paragraph()
    p.text = "衷心感謝：最佳戰友 Carolyn ｜ 推薦人 江沛璇 ｜ 健管教練 麥淑媛 ｜ 寰宇南京中心全體夥伴"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.space_after = Pt(10)

    p = tf11.add_paragraph()
    p.text = "謝謝大家！祝大家都能找回充滿生命力的自己！"
    p.font.name = "Microsoft JhengHei"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_PRIMARY

    slide11.notes_slide.notes_text_frame.text = (
        "【結尾謝幕｜預計 40 秒】\n"
        "最後，我想用我這 45 天刻在心底的一句話做總結：\n"
        "『什麼事情都可以放下，唯獨不能放下筷子；什麼事情都可以失去，唯獨不能失去健康。』\n"
        "健管不是叫我們放下筷子挨餓節食，而是教我們好好拿穩筷子、吃對每一口營養。\n"
        "因為在這個世界上，任何忙碌都可以先放下，唯有這雙筷子下的健康，是我們一輩子絕對不能失去的底氣！\n"
        "48 歲不是衰老的藉口，而是精準管理生命的黃金起點。\n"
        "特別感謝我的戰友 Carolyn、我的推薦人沛璇、麥淑媛教練，以及這段時間互相激勵的所有夥伴。\n"
        "謝謝大家！願我們每一個人都能做自己生命裡的英雄！"
    )

    output_path = "健管45天英雄旅程分享-王執定.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides.")

if __name__ == "__main__":
    create_presentation()
