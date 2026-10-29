import os
import shutil
from PIL import Image, ImageDraw, ImageFont, ImageFilter
import pptx
from pptx.util import Inches

def build_all_nanobanana_deck():
    w, h = 1536, 864
    font_path = "/System/Library/Fonts/STHeiti Medium.ttc"
    
    # 1. 處理 Slide 1 (Nano Banana Pro)
    raw1_path = "/Users/garfiwang/.gemini/antigravity/brain/5a31b131-0cdf-466b-be7c-38835e8014bc/slide1_cover_richman_1793292263807.jpg"
    dst1_raw = "generated/slide01_nanobanana_cover_raw.jpg"
    shutil.copyfile(raw1_path, dst1_raw)
    
    img1 = Image.open(dst1_raw).convert("RGBA")
    rw1, rh1 = img1.size
    target_w = int(rh1 * 16 / 9)
    if rw1 >= target_w:
        cx = (rw1 - target_w) // 2
        c1 = img1.crop((cx, 0, cx + target_w, rh1))
    else:
        th = int(rw1 * 9 / 16)
        cy = (rh1 - th) // 2
        c1 = img1.crop((0, cy, rw1, cy + th))
    b1 = c1.resize((w, h), Image.Resampling.LANCZOS)
    
    # Slide 1 漸層與文字
    grad1 = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    fade_end1 = int(w * 0.58)
    gpix1 = grad1.load()
    for x in range(w):
        if x < fade_end1:
            t = x / fade_end1
            factor = 1.0 - (3 * (t**2) - 2 * (t**3))
            alpha = int(248 * factor)
        else:
            alpha = 0
        for y in range(h):
            gpix1[x, y] = (8, 14, 24, alpha)
    comp1 = Image.alpha_composite(b1, grad1)
    
    t1 = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    s1 = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    dt1 = ImageDraw.Draw(t1)
    ds1 = ImageDraw.Draw(s1)
    
    font_eng = ImageFont.truetype(font_path, 32)
    font_title1 = ImageFont.truetype(font_path, 136)
    font_sub1 = ImageFont.truetype(font_path, 54)
    font_auth = ImageFont.truetype(font_path, 38)
    font_stats = ImageFont.truetype(font_path, 30)
    
    sx, sy = 110, 100
    # 眉標
    ds1.text((sx, sy), "T H E   H E R O ' S   J O U R N E Y", font=font_eng, fill=(0, 0, 0, 220))
    dt1.text((sx, sy), "T H E   H E R O ' S   J O U R N E Y", font=font_eng, fill=(52, 211, 153, 255))
    # 主標
    ds1.text((sx, sy + 60), "45 天的", font=font_title1, fill=(0, 0, 0, 240))
    dt1.text((sx, sy + 60), "45 天的", font=font_title1, fill=(255, 255, 255, 255))
    ds1.text((sx, sy + 215), "蛻變歷程", font=font_title1, fill=(0, 0, 0, 240))
    dt1.text((sx, sy + 215), "蛻變歷程", font=font_title1, fill=(255, 255, 255, 255))
    # 副標
    ds1.text((sx, sy + 400), "重啟 48 歲的身體主控權", font=font_sub1, fill=(0, 0, 0, 240))
    dt1.text((sx, sy + 400), "重啟 48 歲的身體主控權", font=font_sub1, fill=(245, 158, 11, 255))
    # 主講人
    ds1.text((sx, sy + 490), "主講人：王執定（Rich） ｜ 金融男神", font=font_auth, fill=(0, 0, 0, 240))
    dt1.text((sx, sy + 490), "主講人：王執定（Rich） ｜ 金融男神", font=font_auth, fill=(255, 255, 255, 255))
    # 數據
    ds1.text((sx, sy + 555), "45 天累減 5.7kg  |  腰圍大瘦 8cm  |  身體年齡年輕 6 歲", font=font_stats, fill=(0, 0, 0, 240))
    dt1.text((sx, sy + 555), "45 天累減 5.7kg  |  腰圍大瘦 8cm  |  身體年齡年輕 6 歲", font=font_stats, fill=(224, 242, 254, 255))
    
    sb1 = s1.filter(ImageFilter.GaussianBlur(radius=8))
    out1 = Image.alpha_composite(Image.alpha_composite(comp1, sb1), t1)
    slide1_path = "generated/slide01_richman_nanobanana_perfect.png"
    out1.convert("RGB").save(slide1_path, quality=95)
    print("Slide 1 generated:", slide1_path)

    # 2. 處理 Slide 2 (Nano Banana Pro)
    raw2_path = "/Users/garfiwang/.gemini/antigravity/brain/5a31b131-0cdf-466b-be7c-38835e8014bc/slide2_richman_morning_routine_v2_1793292300211.jpg"
    dst2_raw = "generated/slide02_nanobanana_routine_raw.jpg"
    shutil.copyfile(raw2_path, dst2_raw)
    
    img2 = Image.open(dst2_raw).convert("RGBA")
    rw2, rh2 = img2.size
    target_w2 = int(rh2 * 16 / 9)
    if rw2 >= target_w2:
        cx2 = (rw2 - target_w2) // 2
        c2 = img2.crop((cx2, 0, cx2 + target_w2, rh2))
    else:
        th2 = int(rw2 * 9 / 16)
        cy2 = (rh2 - th2) // 2
        c2 = img2.crop((0, cy2, rw2, cy2 + th2))
    b2 = c2.resize((w, h), Image.Resampling.LANCZOS)
    
    comp2 = Image.alpha_composite(b2, grad1)
    t2 = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    s2 = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    dt2 = ImageDraw.Draw(t2)
    ds2 = ImageDraw.Draw(s2)
    
    # 眉標
    ds2.text((sx, sy), "A C T   I   :   T H E   O R D I N A R Y   W O R L D", font=font_eng, fill=(0, 0, 0, 220))
    dt2.text((sx, sy), "A C T   I   :   T H E   O R D I N A R Y   W O R L D", font=font_eng, fill=(52, 211, 153, 255))
    # 主標
    ds2.text((sx, sy + 60), "別人朝九晚五", font=font_title1, fill=(0, 0, 0, 240))
    dt2.text((sx, sy + 60), "別人朝九晚五", font=font_title1, fill=(255, 255, 255, 255))
    ds2.text((sx, sy + 215), "我「朝五晚九」", font=font_title1, fill=(0, 0, 0, 240))
    dt2.text((sx, sy + 215), "我「朝五晚九」", font=font_title1, fill=(255, 255, 255, 255))
    # 副標
    ds2.text((sx, sy + 400), "剛退伍拚到半夜，如今三杯咖啡爬不起來", font=font_sub1, fill=(0, 0, 0, 240))
    dt2.text((sx, sy + 400), "剛退伍拚到半夜，如今三杯咖啡爬不起來", font=font_sub1, fill=(245, 158, 11, 255))
    # 路徑
    ds2.text((sx, sy + 495), "清晨 5:00 BNI 開會  |  整理財經資料  |  下午諮詢演講  |  晚間應酬", font=font_stats, fill=(0, 0, 0, 240))
    dt2.text((sx, sy + 495), "清晨 5:00 BNI 開會  |  整理財經資料  |  下午諮詢演講  |  晚間應酬", font=font_stats, fill=(224, 242, 254, 255))
    
    sb2 = s2.filter(ImageFilter.GaussianBlur(radius=8))
    out2 = Image.alpha_composite(Image.alpha_composite(comp2, sb2), t2)
    slide2_path = "generated/slide02_richman_nanobanana_perfect.png"
    out2.convert("RGB").save(slide2_path, quality=95)
    print("Slide 2 generated:", slide2_path)

    # 3. 建立純血統 Nano Banana Pro 正式版本 PPTX
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    slides_data = [
        (
            slide1_path,
            "【開場｜預計 30 秒】\n大家好，我是 Rich。\n今天要來和大家分享 45 天的蛻變歷程。"
        ),
        (
            slide2_path,
            "【第一幕・平凡世界｜預計 60 秒】\n\n別人都是朝九晚五，而我平常的工作是『朝五晚九』。\n\n什麼叫做朝五晚九呢？大家有聽過 BNI 商會嗎？應該有聽過吧。BNI 是不是都要很早起床？沒有錯。\n\n我有兩個 BNI 商會，一個是我自己創立的長輝分會，另一個是我輔導的金佑分會。所以我平常都很早起來開會，開完會之後要整理財經資料，下午要拜訪客戶或是接受客戶諮詢，有的時候會有演講，晚上也常常會有應酬和飯局。\n\n所以，可能 5 點就要起床，一路忙到晚上九點。\n\n剛退伍的時候，可以拜訪客戶到晚上 12 點，第二天早上 8 點一樣準時進公司開會。\n到了 30 歲以後，早上就會需要喝一杯咖啡來提神。\n到了 40 歲以後，喝一杯咖啡可能不夠了，要喝到三杯咖啡，還會睡過頭爬不起來。"
        ),
        (
            "generated/slide_celebrity_sudden_death_v2.png",
            "【第一幕轉折・新聞震撼｜預計 40 秒】\n\n這時候你就會常常聽到新聞報導很多名人突然猝死。\n\n一個是台灣最大文創旅店承億文旅創辦人戴俊郎，享年才 54 歲！他是大家公認的拼命三郎，心血管疾病吃了好幾年的藥，結果在母親節前夕，突發多發性腦溢血，走了。\n\n另一個更令人震撼：福華飯店第三代廖國宏，才 55 歲！他長期茹素、不抽菸、不喝酒，每天規律健走、打網球，極度自律養生。結果剛下飛機突然抽搐送醫，腦血栓併發腦幹出血，無法開刀，搶救十幾天還是不幸離世。\n\n猝死從來不挑人！身體直接停機登出！\n\n聽到了這些新聞，會不會讓你覺得：『我應該要休息一下？我應該要來做做健管？』"
        ),
        (
            "generated/slide03_richman_crisis_preview.png",
            "【第二幕・冒險召喚｜預計 60 秒】\n\n不會～～因為我覺得我的身體很好呀！我家的基因很健康，我平常都有做好事、累積福報，所以我不會像他們這些人一樣！\n\n後來我的身體給我了一個警告：晚上睡覺半夜，膝蓋突然劇烈痛了起來。\n那種痛，會讓你痛到發瘋！請問大家，你覺得是什麼毛病？\n（等待台下回答：痛風！）\n痛風？！不可能！我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我怎麼可能有痛風？！\n\n後來我的身體，再給了我第二個警告。\n因為我的腰圍越來越寬廣，我就想說應該要來執行 168 斷食，所以我不吃早餐，早上只喝黑咖啡，想說這樣可以燃燒脂肪。\n結果有一天早上，突然頭暈目眩、心跳加速，血壓飆到 160 心跳加速差點昏倒，整個身體完全提不上力，趕快去醫院檢查！\n\n到了這個時候，你覺得是不是應該要讓身體休息一下，要來做做健管，重新讓身體修復了？\n不用啦～～我那麼健康，我家的基因那麼好，我平常都有做好事、積福報，我多喝水，多上健身房運動就可以了啦！"
        ),
        (
            "generated/slide04_threshold_preview.png",
            "【第三幕・跨越門檻｜預計 55 秒】\n\n所以像我這樣鐵齒的人，通常最適合用哪一句成語來形容？\n『不見棺材不掉淚』！\n\n到底是因為什麼事情發生，我才終於肯啟動健管呢？\n關鍵在——城隍爺！對，你沒聽錯，就是掌管地府的城隍爺！\n\n原來是這樣，我每年都會去城隍爺前面『批流年』。\n以往城隍爺給我的提醒都是：『小心開車』、『早點睡覺』。\n但是今年他給我的提醒，竟然是：『小心頭部的疾病』！\n\n請問大家，你如果聽到掌管地府的神明這樣跟你講，你會不會覺得很恐怖？（～會吧～）\n那我就再追問城隍爺：『請問是什麼頭部疾病呀？』\n結果祂就——搖、搖、頭……（學神明搖頭動作）\n\n我看祂一搖頭，我就自己瘋狂腦補：天啊！掌管生死的老闆對著我搖頭，我是不是要中風？然後頭要爆炸？！\n那該怎麼辦？！——『健管！』\n\n所以我馬上請我的引薦人沛璇老師趕快安排健管教練，營養品趕快買一買，立刻執行！"
        )
    ]
    
    for img_p, notes in slides_data:
        s = prs.slides.add_slide(blank_layout)
        s.shapes.add_picture(img_p, 0, 0, width=prs.slide_width, height=prs.slide_height)
        s.notes_slide.notes_text_frame.text = notes
        
    out_pptx = "健管45天圖片簡報-王執定.pptx"
    prs.save(out_pptx)
    print(f"Successfully built 100% Pure Nano Banana Pro deck: {out_pptx} with {len(prs.slides)} slides.")

if __name__ == "__main__":
    build_all_nanobanana_deck()
